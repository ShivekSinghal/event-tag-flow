"""Build printable rules and a QR directory from the React page's exact data."""
import argparse
import json
from pathlib import Path
from xml.sax.saxutils import escape

from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, KeepTogether

PINK = colors.HexColor('#ff007f')
GENERAL = [
    'Play fair. The runner explains the format and decides disputes. No interference; respect consent. Alcohol is always optional, with a non-alcoholic alternative.',
    'Paid play starts after the runner confirms your coin payment. Once a game starts, its entry fee is non-refundable. Free games need no coin payment or NFC scan.',
    'Confirm the fee and local format with the runner before paying. Donations start at 150 whole coins. Party entry is 18+.',
]


def make_pdf(game, output):
    body = ParagraphStyle('Body', fontName='Helvetica', fontSize=10.5, leading=14, spaceAfter=7)
    title = ParagraphStyle('Title', parent=body, fontName='Helvetica-Bold', fontSize=25, leading=29, spaceAfter=10)
    label = ParagraphStyle('Label', parent=body, fontName='Helvetica-Bold', textColor=PINK, fontSize=10, leading=13)
    note = ParagraphStyle('Note', parent=body, fontSize=9, leading=12, textColor=colors.HexColor('#454545'))
    step = ParagraphStyle('Step', parent=body, leftIndent=17, firstLineIndent=-17, alignment=TA_LEFT)
    doc = SimpleDocTemplate(str(output), pagesize=A4, rightMargin=19*mm, leftMargin=19*mm,
                           topMargin=27*mm, bottomMargin=20*mm, title=f"PINK'D {game['name']} Rules",
                           author="PINK'D", invariant=1)

    def frame(canvas, document):
        canvas.saveState()
        width, height = A4
        canvas.setFillColor(colors.black)
        canvas.rect(0, height-19*mm, width, 19*mm, fill=1, stroke=0)
        canvas.setFillColor(colors.white)
        canvas.setFont('Helvetica-Bold', 14)
        canvas.drawString(19*mm, height-12*mm, "PINK'D")
        canvas.setFillColor(PINK)
        canvas.setFont('Helvetica-Bold', 9)
        canvas.drawRightString(width-19*mm, height-12*mm, 'GAME RULES')
        canvas.setFillColor(colors.HexColor('#555555'))
        canvas.setFont('Helvetica', 8)
        canvas.drawString(19*mm, 11*mm, game['url'])
        canvas.linkURL(game['url'], (19*mm, 9*mm, width-30*mm, 15*mm), relative=0)
        canvas.drawRightString(width-19*mm, 11*mm, str(document.page))
        canvas.restoreState()

    story = [Paragraph(escape(f"{game['group']}  |  {game['cost']}"), label),
             Paragraph(escape(game['name']), title), Paragraph(escape(game['outcome']), body)]
    if game.get('prize'):
        story.append(Paragraph(escape('Prize: ' + game['prize']), label))
    story += [Spacer(1, 7), Paragraph('HOW TO PLAY', label)]
    for index, rule in enumerate(game['rules'], 1):
        text = escape(rule) if isinstance(rule, str) else f"<b>{escape(rule['title'])}</b><br/>{escape(rule['body'])}"
        story.append(KeepTogether([Paragraph(f'<b>{index}.</b> {text}', step)]))
    if game.get('note'):
        story += [Spacer(1, 4), Paragraph(escape(game['note']), note)]
    story += [Spacer(1, 9), Paragraph('BEFORE YOU PLAY', label)]
    for text in GENERAL:
        story.append(Paragraph(escape(text), note))
    if game.get('prize'):
        story.append(Paragraph('Each Pinkredible gives Rs. 100 off course registration, not cash. Check your band on My coins.', note))
    doc.build(story, onFirstPage=frame, onLaterPages=frame)


def make_docx(games, assets, output):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(.6)
    section.left_margin = section.right_margin = Inches(.75)
    for name in ['Normal', 'Title', 'Heading 1']:
        style = doc.styles[name]
        style.font.name = 'Arial'
        style.font.color.rgb = RGBColor(0, 0, 0)
    doc.styles['Normal'].font.size = Pt(10)
    doc.styles['Normal'].paragraph_format.space_after = Pt(6)
    doc.styles['Title'].font.size = Pt(23)
    doc.styles['Heading 1'].font.size = Pt(16)
    doc.core_properties.title = 'Pinkd individual game QR codes'
    doc.core_properties.author = "PINK'D"

    def link(label, url):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        hyperlink = OxmlElement('w:hyperlink')
        hyperlink.set(qn('r:id'), p.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
        run = OxmlElement('w:r')
        props = OxmlElement('w:rPr')
        color = OxmlElement('w:color')
        color.set(qn('w:val'), 'C00060')
        props.append(color)
        underline = OxmlElement('w:u')
        underline.set(qn('w:val'), 'single')
        props.append(underline)
        run.append(props)
        text = OxmlElement('w:t')
        text.text = label
        run.append(text)
        hyperlink.append(run)
        p._p.append(hyperlink)

    for i, game in enumerate(games):
        if i % 2 == 0:
            if i:
                doc.add_page_break()
            p = doc.add_paragraph('Pinkd game QR codes', 'Title')
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            intro = 'Scan a code to open that game\'s rules and download its PDF. Print at full size and keep the white border around each QR code.'
            if i:
                intro = 'Scan to read the rules and download the individual game PDF.'
            p = doc.add_paragraph(intro)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p = doc.add_paragraph(f"{i+1}  {game['name']}", 'Heading 1')
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(15)
        p = doc.add_paragraph(f"{game['group']}  |  {game['cost']}")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture = p.add_run().add_picture(str(assets / f"{game['routeId']}.png"), width=Inches(2))
        picture._inline.docPr.set('descr', f"QR code for {game['name']}: {game['url']}")
        link(game['url'], game['url'])
        link('Download this game PDF', f"https://pinkd.hashtag.dance/game-rules/{game['routeId']}.pdf")
    # Some Word base templates carry a blue title border; keep the QR sheet plain.
    for root in [doc.styles.element, doc.element]:
        for border in root.iter(qn('w:pBdr')):
            border.getparent().remove(border)
    doc.save(output)


parser = argparse.ArgumentParser()
parser.add_argument('data', type=Path)
parser.add_argument('assets', type=Path)
parser.add_argument('--docx', type=Path)
args = parser.parse_args()
games = json.loads(args.data.read_text())
for game in games:
    make_pdf(game, args.assets / f"{game['routeId']}.pdf")
if args.docx:
    make_docx(games, args.assets, args.docx)
